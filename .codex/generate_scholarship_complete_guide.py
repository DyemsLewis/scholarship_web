from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(
    r"D:\XAMPP\htdocs\scholarship_web\deliverables\Scholarships_Complete_Guide_and_Applicant_Processes.docx"
)

INK = "111827"
NAVY = "17324D"
BLUE = "2E5F8A"
SLATE = "526273"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "D9E0E7"
WHITE = "FFFFFF"


def set_font(run, name="Aptos", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text, bold=False, color=INK, size=9.3):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(str(text))
    set_font(run, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_column_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def add_table(document, headers, rows, widths=None, font_size=9.1):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    set_row_cant_split(header_row)
    for idx, header in enumerate(headers):
        shade_cell(header_row.cells[idx], NAVY)
        set_cell_text(header_row.cells[idx], header, bold=True, color=WHITE, size=9.1)
    for row_data in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for idx, value in enumerate(row_data):
            if len(table.rows) % 2 == 1:
                shade_cell(row.cells[idx], LIGHT_GRAY)
            set_cell_text(row.cells[idx], value, size=font_size)
    if widths:
        set_column_widths(table, widths)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([run_properties, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)
    section.different_first_page_header_footer = True

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.13
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(INK)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    subtitle.font.size = Pt(12)
    subtitle.font.color.rgb = RGBColor.from_string(SLATE)
    subtitle.paragraph_format.space_after = Pt(16)

    for style_name, size, color, before, after in [
        ("Heading 1", 18, NAVY, 18, 8),
        ("Heading 2", 14, BLUE, 14, 6),
        ("Heading 3", 11.5, NAVY, 10, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(10.2)
        style.paragraph_format.left_indent = Inches(0.29)
        style.paragraph_format.first_line_indent = Inches(-0.17)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.widow_control = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("SCHOLARSHIPS AND APPLICANT PROCESSES")
    set_font(run, size=8, color=SLATE, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Complete reference guide  |  ")
    set_font(run, size=8, color=SLATE)
    add_page_number(footer)
    for footer_run in footer.runs:
        set_font(footer_run, size=8, color=SLATE)

    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_callout(document, label, text, fill=LIGHT_BLUE):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.75)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=130, start=150, bottom=130, end=150)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(f"{label}: ")
    set_font(run, size=9.8, color=NAVY, bold=True)
    run = paragraph.add_run(text)
    set_font(run, size=9.8, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_labeled_paragraph(document, label, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(f"{label}: ")
    set_font(run, bold=True, color=NAVY)
    paragraph.add_run(text)


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):
            label, detail = item
            run = paragraph.add_run(label)
            set_font(run, bold=True, color=NAVY)
            paragraph.add_run(detail)
        else:
            paragraph.add_run(item)


def add_numbered(document, items):
    for index, item in enumerate(items, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.35)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.1
        paragraph.paragraph_format.widow_control = True
        number_run = paragraph.add_run(f"{index}. ")
        set_font(number_run, bold=True, color=NAVY)
        if isinstance(item, tuple):
            label, detail = item
            run = paragraph.add_run(label)
            set_font(run, bold=True, color=NAVY)
            paragraph.add_run(detail)
        else:
            paragraph.add_run(item)


def add_contents(document):
    document.add_heading("Contents", level=1)
    items = [
        "Part I. Understanding scholarships and educational assistance",
        "Part II. Scholarship types, benefits, and program design",
        "Part III. Eligibility, evidence, and readiness",
        "Part IV. The complete scholarship process",
        "Part V. Processes for each applicant type",
        "Part VI. Selection, provider operations, and decisions",
        "Part VII. Privacy, safety, fairness, and problem handling",
        "Part VIII. Practical checklists, glossary, and official references",
    ]
    for item in items:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        set_font(run, bold=True, color=NAVY)
    document.add_page_break()


def add_applicant_process(
    document,
    title,
    who_leads,
    use_case,
    core_data,
    evidence,
    steps,
    safeguards,
    blockers,
):
    document.add_heading(title, level=2)
    add_labeled_paragraph(document, "Who normally leads", who_leads)
    add_labeled_paragraph(document, "What makes this applicant type different", use_case)
    document.add_heading("Information normally needed", level=3)
    add_bullets(document, core_data)
    document.add_heading("Evidence commonly requested", level=3)
    add_bullets(document, evidence)
    document.add_heading("Recommended process", level=3)
    add_numbered(document, steps)
    document.add_heading("Safeguards and practical adjustments", level=3)
    add_bullets(document, safeguards)
    document.add_heading("Common reasons the application may stop or require correction", level=3)
    add_bullets(document, blockers)


def build_document():
    document = Document()
    configure_document(document)

    title = document.add_paragraph(style="Title")
    title.add_run("Scholarships and Applicant Processes")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run(
        "A complete guide to what scholarships are, how programs differ, what evidence is needed, and how the process changes for each learner group"
    )

    lead = document.add_paragraph()
    lead.paragraph_format.space_after = Pt(10)
    run = lead.add_run(
        "A scholarship is not a single, standard product. It is a defined package of educational support with a purpose, target group, benefit, eligibility rules, evidence requirements, selection method, release conditions, and continuing obligations. A fair scholarship process therefore changes according to the learner's education level, age, legal capacity, circumstances, and intended learning pathway."
    )
    set_font(run, size=11.2, color=INK, bold=True)

    document.add_paragraph(
        "This guide is written for Filipino applicants, parents or guardians, scholarship providers, schools, reviewers, and designers of scholarship-finder systems. It explains common practice and uses official Philippine programs as examples. It does not replace the current guidelines of CHED, DepEd, TESDA, DOST-SEI, an LGU, a school, or a private provider. When a program announcement differs from this guide, the official and current program rules control."
    )
    add_callout(
        document,
        "Central conclusion",
        "Eligibility is only the first gate. Meeting minimum criteria means an application may proceed to review; it does not guarantee a slot, award, admission, payment, or final acceptance. Competitive programs may rank qualified applicants because funding and slots are limited.",
    )
    add_callout(
        document,
        "Platform context",
        "A scholarship finder and eligibility platform should act as a discovery and pre-screening channel. It can organize applicant data, identify likely matches, collect the minimum evidence needed for review, and communicate next steps. The verified provider still makes the final decision and may require originals or a separate formal application later.",
        fill=LIGHT_GRAY,
    )
    document.add_paragraph("Prepared as a comprehensive educational and operational reference. Last reviewed: September 2026.")
    document.add_page_break()
    add_contents(document)

    document.add_heading("Part I. Understanding Scholarships and Educational Assistance", level=1)

    document.add_heading("1. What a scholarship is", level=2)
    document.add_paragraph(
        "A scholarship is educational assistance awarded to a learner under published or documented conditions. The support is intended to help the learner enter, continue, or complete education or training. It may recognize merit, respond to financial need, develop talent, support a priority field, serve a community, or combine several purposes."
    )
    document.add_paragraph(
        "The word scholarship is often used broadly, but the actual arrangement matters more than the label. One program may pay tuition directly to a school, another may release a living allowance, another may provide a laptop or review classes, and another may offer mentoring plus an internship without paying tuition. Applicants should read the benefit schedule, release method, renewal rules, and obligations instead of assuming that every scholarship is a full cash award."
    )

    document.add_heading("What a scholarship normally contains", level=3)
    add_bullets(document, [
        ("A defined purpose. ", "The provider explains the educational or social result the fund is intended to support."),
        ("A target applicant group. ", "The program identifies the education levels, locations, fields, circumstances, or communities it intends to serve."),
        ("A support package. ", "The provider states what is covered, the amount or quantity, the frequency, the release method, and any exclusions."),
        ("Eligibility rules. ", "These are the minimum conditions an applicant must meet before competitive review."),
        ("Evidence requirements. ", "Documents or verifiable records support claims about enrollment, grades, income, location, identity, or special circumstances."),
        ("A selection process. ", "The provider may use completeness review, a rubric, ranking, an exam, an interview, portfolio review, or a combination."),
        ("Decision and communication rules. ", "The program should state who decides, how applicants are notified, whether there is a waitlist, and how corrections or concerns are handled."),
        ("Release and continuing conditions. ", "The recipient may need to enroll, maintain academic standing, submit reports, attend activities, or comply with a service commitment."),
    ])

    document.add_heading("What a scholarship does not automatically mean", level=3)
    add_bullets(document, [
        "It does not automatically guarantee admission to a school, course, or training institution unless the program expressly includes admission.",
        "It does not necessarily cover every school expense or continue for the entire program.",
        "It is not automatically a cash payment to the learner; the provider may pay the institution or reimburse approved expenses.",
        "Passing a platform pre-screen does not mean the provider has granted the award.",
        "Being academically strong does not override location, income, course, citizenship, age, or other stated restrictions.",
        "A provider may withdraw an offer when material information is false, required conditions are not met, or the applicant does not accept by the deadline, subject to the program rules and fair notice.",
    ])

    document.add_heading("2. Why scholarships exist", level=2)
    document.add_paragraph(
        "Scholarships convert public, institutional, philanthropic, or private resources into educational opportunity. Providers do not all pursue the same goal, so two programs may reasonably ask for different evidence and use different selection methods. The provider should be able to connect each eligibility rule and document request to a legitimate program purpose."
    )
    add_table(
        document,
        ["Program purpose", "What it tries to achieve", "Typical selection emphasis"],
        [
            ("Educational access", "Reduce financial barriers that prevent enrollment or continuation.", "Need, enrollment status, cost, and ability to continue."),
            ("Academic merit", "Recognize strong academic performance or potential.", "Grades, rank, assessment, achievements, and consistency."),
            ("Equity and inclusion", "Improve access for underserved or underrepresented learners.", "Relevant circumstance plus readiness and program fit."),
            ("Priority workforce fields", "Build skills in fields needed by society, industry, science, or public service.", "Chosen field, aptitude, commitment, and completion potential."),
            ("Talent development", "Develop ability in arts, sports, leadership, innovation, or research.", "Portfolio, audition, competition record, proposal, or demonstrated potential."),
            ("Community development", "Support residents or learners who can contribute to a locality or community.", "Residence, community connection, service record, and intended impact."),
            ("Retention and completion", "Prevent capable learners from stopping because of temporary or continuing barriers.", "Current standing, remaining study period, need, and completion plan."),
            ("Institutional mission", "Advance the goals of a school, foundation, company, profession, or donor.", "Fit with the stated mission and transparent criteria."),
        ],
        widths=[Inches(1.45), Inches(2.6), Inches(2.7)],
    )

    document.add_heading("3. Scholarship compared with related assistance", level=2)
    add_table(
        document,
        ["Form of support", "Main idea", "Important distinction"],
        [
            ("Scholarship", "Educational support awarded using stated eligibility and selection conditions.", "May be merit-based, need-based, targeted, or mixed; may include non-cash benefits."),
            ("Grant-in-aid or bursary", "Assistance commonly driven by financial need or a defined educational expense.", "May use less competitive merit ranking, but still has eligibility and documentation rules."),
            ("Voucher or subsidy", "A public or institutional contribution toward an approved education service.", "Often paid or credited to a participating institution and may not be cash given to the learner."),
            ("Tuition waiver or discount", "The institution reduces all or part of its own tuition or fees.", "Usually usable only at that institution and may not cover living costs."),
            ("Fellowship", "Support for advanced study, research, professional development, or a structured cohort experience.", "Often includes research outputs, milestones, mentoring, or service expectations."),
            ("Sponsorship", "An organization supports a learner, event, training, or educational pathway for a defined objective.", "May include branding, employment pathways, service, or reporting conditions."),
            ("Student loan", "Money must generally be repaid under agreed terms.", "It is not a scholarship even when repayment is deferred or subsidized."),
            ("Prize or award", "Recognition given after an achievement or competition.", "May not fund continuing education and may be one-time."),
            ("Work-study or employment assistance", "The learner earns compensation or support through approved work.", "Work conditions, labor protections, and time commitments matter."),
        ],
        widths=[Inches(1.4), Inches(2.55), Inches(2.8)],
    )
    add_callout(
        document,
        "Philippine policy example",
        "Republic Act No. 10931 treats free higher education, tertiary education subsidies, and student loan programs as related but distinct forms of student financial assistance. Applicants should not assume that every government education benefit is a scholarship.",
    )

    document.add_heading("Part II. Scholarship Types, Benefits, and Program Design", level=1)

    document.add_heading("4. Types based on how recipients are selected", level=2)
    add_bullets(document, [
        ("Merit-based. ", "Prioritizes academic achievement, demonstrated ability, potential, or a competitive assessment. Financial need may be irrelevant, secondary, or used as a tie-breaker."),
        ("Need-based. ", "Prioritizes applicants whose household resources are insufficient for the educational cost. Academic standing may still be required to show readiness."),
        ("Merit-and-need. ", "Requires both academic readiness and demonstrated financial need. A weighted rubric may balance the two."),
        ("Talent-based. ", "Supports ability in sports, music, visual arts, performing arts, writing, technology, invention, or another defined talent. A portfolio, audition, trial, or competition record may be needed."),
        ("Leadership and service. ", "Recognizes sustained leadership, volunteering, civic participation, or community contribution. Review should focus on verified quality and impact, not only the number of activities."),
        ("Field- or course-specific. ", "Supports study in a priority discipline, occupation, research area, or qualification. Changing course may require approval because it changes program fit."),
        ("Location- or community-based. ", "Limits support to residents, students of a locality, indigenous or geographically isolated communities, or another defined coverage area. The rule should state whether residence is mandatory or only preferred."),
        ("Equity-targeted. ", "Supports a population that faces documented barriers, such as learners with disabilities, displaced learners, single-parent households, or other groups named by the program. Only necessary evidence should be collected."),
        ("Dependent, employee, member, or alumni-linked. ", "Eligibility comes from a documented relationship with an employee, cooperative, professional organization, union, school, or alumni association."),
        ("Research or innovation. ", "Supports a proposal, thesis, dissertation, laboratory work, publication, prototype, or knowledge output. Selection emphasizes feasibility, ethics, supervision, and relevance."),
        ("Emergency or continuity assistance. ", "Responds to sudden hardship, disaster, displacement, illness, or another event that threatens continuation. Processing may be faster and evidence proportionate to the emergency."),
        ("Open competitive. ", "Accepts applicants across a broad group and ranks them using a disclosed process. 'Open' does not mean that everyone receives support."),
    ])

    document.add_heading("5. Types based on provider", level=2)
    add_bullets(document, [
        ("National government. ", "Programs may advance education access, science and technology, workforce development, or another national priority and follow formal appropriations and guidelines."),
        ("Local government unit. ", "Programs commonly prioritize residents and may use barangay, city, municipal, or provincial evidence."),
        ("School or training institution. ", "Awards may be tied to admission, tuition, school service, athletics, arts, academic ranking, or institutional need assessment."),
        ("Corporate or employer. ", "Support may be part of social investment, employee benefits, workforce development, or a talent pipeline. Applicants should read internship, employment, or return-service conditions carefully."),
        ("Foundation, nonprofit, religious, or community organization. ", "Programs are shaped by the organization's mission, donor restrictions, service area, and available funds."),
        ("Professional association or cooperative. ", "Support may target a profession, members and dependents, or community development objectives."),
        ("International organization or foreign institution. ", "Programs may involve mobility, visas, language evidence, international admission, insurance, and cross-border data processing."),
        ("Individual or family donor. ", "The program should still use documented criteria, secure data handling, accountable disbursement, and a clear contact channel."),
    ])

    document.add_heading("6. Types based on education pathway", level=2)
    add_table(
        document,
        ["Pathway", "Typical purpose", "Important process difference"],
        [
            ("Basic education", "Access, supplies, meals, transport, tuition assistance, or continuity.", "A parent or guardian usually leads for younger minors; data and documents should be minimized."),
            ("Senior high school", "Tuition support, strand development, transition, equipment, or career preparation.", "Track/strand, school participation, voucher status, and guardian involvement may matter."),
            ("Alternative Learning System", "Completion, equivalency, transition, or re-entry support.", "ALS enrollment, portfolio, assessment or equivalency evidence may replace conventional year-level records."),
            ("TVET", "Training, assessment, certification, tools, transport, or employment preparation.", "Qualification, training provider, schedule, assessment, and certification are central."),
            ("Undergraduate", "Tuition, allowance, course support, retention, and completion.", "Admission/enrollment, course, units, grades, income, ranking, and renewal are common."),
            ("Graduate and postgraduate", "Advanced professional study, research, faculty development, or public-service capacity.", "Proposal quality, adviser, research ethics, milestones, publications, or return service may apply."),
            ("Short course or continuing education", "Specific skills, licensure preparation, upskilling, or reskilling.", "Course legitimacy, schedule, completion evidence, and job relevance matter more than a conventional grade level."),
        ],
        widths=[Inches(1.35), Inches(2.55), Inches(2.85)],
    )

    document.add_heading("7. Types based on coverage and duration", level=2)
    add_bullets(document, [
        ("Full scholarship. ", "Covers all items expressly listed by the provider; it still may exclude optional fees, penalties, devices, housing, or other costs."),
        ("Partial scholarship. ", "Pays a percentage, fixed amount, or selected cost categories, leaving a funding gap the learner must plan for."),
        ("Tuition-only. ", "Covers tuition or an approved portion of it but not necessarily fees, transport, food, lodging, supplies, or devices."),
        ("Allowance or stipend. ", "Provides recurring support for living or study costs, usually subject to enrollment and reporting."),
        ("In-kind support. ", "Provides goods or services such as uniforms, devices, connectivity, books, review classes, mentoring, or dormitory space."),
        ("One-time award. ", "Provides support once for a semester, school year, project, examination, emergency, or defined purchase."),
        ("Renewable award. ", "May continue for multiple terms when the recipient meets renewal conditions and funds remain available."),
        ("Reimbursement. ", "Pays only after the recipient submits approved receipts or evidence of an eligible expense."),
        ("Direct institutional payment. ", "The provider pays or credits the school or training institution rather than transferring the whole benefit to the learner."),
    ])

    document.add_heading("8. Benefits a scholarship may provide", level=2)
    document.add_paragraph(
        "A complete program description should list every benefit separately. Combining all support into one 'award amount' can mislead applicants because many valuable benefits are not cash and may follow different release schedules."
    )
    add_table(
        document,
        ["Benefit category", "Examples", "What the provider should disclose"],
        [
            ("Tuition and school charges", "Tuition, laboratory, miscellaneous, assessment, graduation, or approved institutional fees.", "Maximum amount, covered fee types, participating institutions, and payment recipient."),
            ("Living and study allowance", "Monthly or term allowance, meals, transport, boarding, or dormitory support.", "Amount, frequency, first release date, conditions, and whether liquidation is required."),
            ("Learning materials", "Books, uniform, supplies, tools, equipment, device, data, or connectivity.", "Exact item or cap, ownership, replacement, and return rules."),
            ("Training and assessment", "Review classes, technical training, competency assessment, certification, or licensure preparation.", "Approved provider, schedule, fees, attendance, and completion requirements."),
            ("Academic and personal support", "Tutoring, mentoring, coaching, counseling, accommodations, or advising.", "Availability, confidentiality, participation expectations, and referral process."),
            ("Career support", "Internship, job exposure, networking, career guidance, placement referral, or employer mentoring.", "Whether participation is guaranteed, optional, competitive, paid, or tied to service."),
            ("Research support", "Thesis costs, data collection, laboratory use, conference travel, publication, or equipment.", "Allowable expenses, ethics approval, output ownership, reporting, and liquidation."),
            ("Health and protection", "Insurance, medical allowance, emergency support, or disability accommodations.", "Scope, exclusions, claims process, privacy handling, and responsible party."),
        ],
        widths=[Inches(1.4), Inches(2.55), Inches(2.8)],
    )

    document.add_heading("9. What a complete program announcement should show", level=2)
    add_bullets(document, [
        "Verified provider name, public contact details, and a safe method for confirming authenticity.",
        "Program title, purpose, target applicant group, number of slots or a statement that slots are limited.",
        "Opening date, deadline, time zone, and whether late or incomplete submissions are accepted.",
        "Exact eligibility rules, including whether a condition is mandatory, preferred, or only used for ranking.",
        "Complete benefit list, amounts or caps, frequency, duration, release method, exclusions, and tax or bank requirements if applicable.",
        "Required and supporting documents, acceptable file formats, validity periods, and whether originals are requested later.",
        "Selection stages, criteria or rubric, exam/interview mode, likely schedule window, and who makes the final decision.",
        "How applicants may correct information, withdraw, report a concern, request an accommodation, or ask about privacy.",
        "Renewal, suspension, termination, appeal, service, attendance, reporting, and change-of-school/course rules.",
    ])

    document.add_heading("Part III. Eligibility, Evidence, and Readiness", level=1)

    document.add_heading("10. Understanding eligibility", level=2)
    document.add_paragraph(
        "Eligibility is a set of minimum conditions. A condition should be measurable, relevant to the program purpose, stated before submission, and applied consistently. Providers should avoid hidden rules and should not collect a document merely because it has traditionally been requested."
    )
    add_table(
        document,
        ["Eligibility area", "Examples", "Questions that must be clear"],
        [
            ("Identity or status", "Citizenship, residency, age, applicant relationship, or membership.", "Is it mandatory? At what date is it measured? What evidence is accepted?"),
            ("Education", "Education level, grade/year, ALS status, TVET qualification, school type, course, track, strand, or units.", "Are incoming and current learners both eligible? Are transferees or returnees included?"),
            ("Academic", "General average, GWA/GPA, grade-point threshold, no failing grade, rank, or assessment score.", "Which grading scale and period are used? Is conversion allowed? Is there an alternative when no grades exist?"),
            ("Financial", "Household income, per-capita income, need category, or special hardship.", "What period and household definition apply? Which proofs or alternatives are accepted?"),
            ("Location", "Barangay, city, province, region, school location, or distance from a service site.", "Is coverage strict, preferred, or only used to calculate access?"),
            ("Field and pathway", "Priority course, occupation, research area, strand, qualification, or institution.", "Are equivalent course names accepted? Can the learner change path later?"),
            ("Other assistance", "Restrictions on holding another scholarship or receiving overlapping benefits.", "Which combinations are prohibited, and must existing support be disclosed?"),
            ("Participation", "Availability for an exam, interview, orientation, mentoring, internship, or service.", "Is attendance mandatory, and will an accommodation or alternate schedule be considered?"),
        ],
        widths=[Inches(1.25), Inches(2.6), Inches(2.9)],
    )

    document.add_heading("Eligible, matched, shortlisted, and selected are different", level=3)
    add_bullets(document, [
        ("Potential match. ", "The profile appears compatible with available program data, but required fields or evidence may still be missing."),
        ("Eligible. ", "The applicant appears to satisfy every mandatory rule based on available information."),
        ("Document-complete. ", "All files required for the current stage were submitted in an acceptable form; authenticity may still need verification."),
        ("Verified. ", "An authorized reviewer checked the relevant claims or evidence. Verification should specify what was checked and by whom."),
        ("Shortlisted. ", "The applicant remains under consideration after initial review or ranking."),
        ("Qualified for the next stage. ", "The applicant may proceed to an exam, interview, formal application, or other provider step."),
        ("Selected or awarded. ", "The authorized provider made a final offer, normally subject to acceptance and any final conditions."),
        ("Waitlisted. ", "The applicant is eligible but no current slot is available; selection may occur only if a slot opens."),
    ])

    document.add_heading("11. Academic measures and grading systems", level=2)
    document.add_paragraph(
        "Programs should not assume that every applicant uses the same grading scale. Philippine institutions may use percentages, GWA, GPA, grade points, competency results, descriptive marks, or another approved system. The provider should record the original scale and avoid an undocumented conversion."
    )
    add_bullets(document, [
        "Ask for the grading scale and period covered, such as the latest completed grading period, semester, or school year.",
        "Allow the provider to define a percentage threshold, GWA/GPA threshold, grade-point rule, competency result, rank, or no academic minimum.",
        "If conversion is needed, publish the conversion method or request institutional certification; do not let applicants guess.",
        "For kindergarten, early elementary, ALS, competency-based TVET, or programs without conventional numeric grades, use age-appropriate or pathway-appropriate evidence instead of forcing a GWA.",
        "Treat OCR or automated extraction as an aid. A reviewer should compare extracted grades with the source image or official record before verification.",
    ])

    document.add_heading("12. Applicant information and evidence", level=2)
    document.add_paragraph(
        "The correct question is not 'What documents can we collect?' but 'What claim must be checked at this stage, and what is the least intrusive reliable evidence?' A pre-screening platform should collect reusable, low-risk information first and defer sensitive originals or provider-specific forms until they are genuinely needed."
    )
    add_table(
        document,
        ["Claim to establish", "Typical evidence", "Less intrusive or alternate evidence"],
        [
            ("Current education status", "Certificate of enrollment, registration form, school ID, or school certification.", "Portal enrollment record or authorized school confirmation."),
            ("Academic standing", "Report card, transcript, certificate of grades, or official grade portal copy.", "School-certified summary or a reviewer-verified extract."),
            ("Household financial condition", "Tax return, certificate of indigency, payslip, employment certificate, or social welfare record.", "Provider-defined affidavit or alternate proof when formal income records do not exist."),
            ("Residence or coverage", "Barangay certificate, government record, billing address, or school record.", "Geocoded address plus later verification when strict residence is essential."),
            ("Identity and age", "Government or school record appropriate to the program.", "A guardian-attested account for early pre-screening, with formal proof deferred when allowed."),
            ("Special category", "A certificate or official record relevant to the stated category.", "A narrowly scoped attestation followed by verification only for shortlisted applicants."),
            ("Relationship or guardianship", "Birth, guardianship, authorization, or school record where necessary.", "Guardian declaration and contact verification for low-risk initial steps."),
        ],
        widths=[Inches(1.4), Inches(2.55), Inches(2.8)],
    )

    document.add_heading("Stage-based document collection", level=3)
    add_numbered(document, [
        ("Discovery and matching. ", "Use profile attributes and applicant-entered information; do not require sensitive documents simply to browse."),
        ("Pre-screen submission. ", "Collect only the files needed to establish mandatory eligibility and application readiness."),
        ("Provider review. ", "Allow targeted corrections and reviewer verification; keep a record of who changed a status and why."),
        ("Shortlist or next stage. ", "Request provider-specific forms, clearer copies, or additional proof only when relevant to the next decision."),
        ("Formal application or award. ", "The provider may inspect originals, request certified copies, validate admission or enrollment, and collect payment details through secure channels."),
        ("Renewal and monitoring. ", "Collect only current records needed to determine continuing eligibility and release the next benefit."),
    ])

    document.add_heading("13. Readiness before applying", level=2)
    add_bullets(document, [
        "Use an email address and contact number the applicant or authorized guardian can access throughout the cycle.",
        "Complete the profile truthfully and keep names, dates, school information, and addresses consistent with supporting records.",
        "Prepare readable files, use descriptive filenames, check expiration dates, and keep originals secure.",
        "Read the provider identity, benefit details, eligibility rules, deadline, process stages, privacy notice, and obligations before agreeing.",
        "Confirm whether another scholarship can be held at the same time and disclose existing support when requested.",
        "Plan for costs not covered by the award and for attendance at any mandatory activity.",
        "Request an accommodation early if disability, connectivity, distance, work, caregiving, or another barrier affects participation.",
    ])

    document.add_heading("Part IV. The Complete Scholarship Process", level=1)

    document.add_heading("14. Universal end-to-end process", level=2)
    document.add_paragraph(
        "Not every program uses every stage, but a complete process should define what happens before, during, and after selection. The following model separates platform pre-screening from the provider's formal decision."
    )
    add_numbered(document, [
        ("Program design and authorization. ", "The provider establishes the goal, budget, slots, target group, benefits, criteria, evidence, rubric, schedule, reviewer authority, privacy basis, and complaint process."),
        ("Publication and authenticity check. ", "The program is published through an official or verified channel. Applicants confirm the provider, domain, contact, deadline, and absence of suspicious fees or requests."),
        ("Discovery and profile matching. ", "The learner browses programs and uses a profile to identify likely matches. Matching explains which criteria appear met, unmet, unknown, or not applicable."),
        ("Eligibility self-check. ", "The learner reviews mandatory conditions before spending time on documents. 'Any' must behave as unrestricted, not as a mismatch."),
        ("Account, authority, and consent. ", "The correct person creates or supports the account. A guardian acts where appropriate for a minor, and privacy and application terms are presented clearly."),
        ("Application preparation. ", "The applicant completes program questions, chooses reusable documents, uploads program-specific evidence, and checks benefit and process details."),
        ("Submission and receipt. ", "The system timestamps the application, locks or versions submitted data, provides a reference number, and shows how corrections and withdrawal work."),
        ("Completeness review. ", "The provider checks that required fields and files are present, readable, current, and relevant. Correctable issues should be returned with a specific explanation and deadline."),
        ("Eligibility and evidence review. ", "Authorized reviewers apply mandatory rules consistently and examine relevant source files. A platform score guides review but does not replace provider judgment."),
        ("Competitive assessment. ", "When qualified applicants exceed available slots, the provider applies a documented rubric, ranking, exam, interview, portfolio, or other announced method."),
        ("Pre-screen outcome. ", "The provider rejects, requests correction, waitlists, or qualifies the applicant for a formal application or the next selection stage."),
        ("Formal provider process. ", "The provider may inspect originals, issue its own official form, validate data with a school, conduct an external exam/interview, and complete due diligence outside the portal."),
        ("Final decision and offer. ", "The authorized decision-maker confirms selected, waitlisted, and unsuccessful applicants. The offer states benefits, conditions, acceptance deadline, and next action."),
        ("Acceptance and onboarding. ", "The selected applicant accepts the award terms, submits only final required records, attends orientation if needed, and provides secure disbursement details."),
        ("Benefit release. ", "The provider pays the institution, transfers an allowance, distributes goods, or delivers services according to the published schedule and records the release."),
        ("Monitoring and renewal. ", "The recipient submits current enrollment, grades, outputs, or other relevant evidence. The provider applies published renewal and appeal rules."),
        ("Completion, transition, or service. ", "The program records completion, graduation, certification, employment transition, research output, alumni engagement, or any lawful agreed service."),
        ("Closure and retention. ", "The provider closes the case, communicates final status, retains records only as required, and securely disposes of data when the legitimate purpose and retention period end."),
    ])

    document.add_heading("15. Corrections, withdrawal, and resubmission", level=2)
    add_bullets(document, [
        ("Before submission. ", "Applicants should be able to edit draft data and replace files freely."),
        ("After submission but before review. ", "The system may permit withdrawal and resubmission or a controlled edit that creates a new version."),
        ("During review. ", "The provider should request a specific correction, identify the affected field or file, and set a reasonable deadline. The original review record should remain auditable."),
        ("After qualification or selection. ", "Material changes should trigger re-evaluation when they affect eligibility, ranking, or identity. Minor contact corrections should not automatically cancel an application."),
        ("Withdrawal. ", "An applicant should be able to withdraw voluntarily, understand whether the action is reversible, and receive confirmation. The provider should not continue processing beyond legitimate retention needs."),
        ("Provider cancellation or amendment. ", "If funding, dates, benefits, or criteria change, affected applicants should receive clear notice and a fair opportunity to respond or withdraw."),
    ])

    document.add_heading("16. Status language applicants can understand", level=2)
    add_table(
        document,
        ["Status", "What it should mean", "Applicant action"],
        [
            ("Draft", "The application has not been submitted.", "Complete required fields and review before the deadline."),
            ("Submitted", "The platform received the application.", "Keep contact details active; no repeated submission is needed."),
            ("Needs correction", "A specific correctable issue prevents review.", "Open the request, replace or clarify the item, and resubmit by the stated date."),
            ("Under review", "The provider is checking eligibility, evidence, and fit.", "Wait for an official update unless information is requested."),
            ("Qualified for next stage", "The pre-screen was passed and another provider step is required.", "Read the schedule and instructions; passing this stage is not yet a final award."),
            ("Waitlisted", "The application is acceptable but no current slot is available.", "Follow the waitlist deadline and do not assume selection."),
            ("Selected", "The provider made an offer subject to the stated final conditions.", "Accept or decline by the deadline and complete onboarding."),
            ("Not selected", "The application did not receive a slot in this cycle.", "Read any reason or appeal route and consider other programs."),
            ("Withdrawn", "The applicant ended the application or an authorized withdrawal was recorded.", "Contact support only if the withdrawal was unauthorized or reversible under policy."),
            ("Completed", "The award cycle or recipient obligation was closed.", "Keep personal records and complete any final report that was disclosed."),
        ],
        widths=[Inches(1.25), Inches(3.0), Inches(2.5)],
    )

    document.add_heading("Part V. Processes for Each Applicant Type", level=1)
    document.add_paragraph(
        "In this guide, applicant type means the learner's educational pathway or circumstances. It is not a label that reduces the applicant to one characteristic. A learner may belong to more than one type, such as a working undergraduate with a disability or an ALS completer applying for TVET. The program should apply the process that best fits the actual pathway and legal context."
    )

    add_applicant_process(
        document,
        "17. Kindergarten and elementary learners",
        "A parent, legal guardian, or properly authorized adult normally creates and manages the application. The child should receive an age-appropriate explanation and should not be expected to manage legal consent or complex uploads.",
        "These applicants are young minors. Academic merit should not be measured only through high-stakes numeric grades, and the process should minimize collection of identity, family, and school records.",
        [
            "Learner name, age or date of birth, current or intended grade level, school, and basic location information.",
            "Parent/guardian identity, relationship, authority, contact details, and preferred communication channel.",
            "Household or access circumstances only when directly relevant to need-based eligibility.",
            "Support needs such as transport, supplies, meals, tuition, device access, or accommodation.",
        ],
        [
            "Enrollment, admission, school certification, or report card appropriate to the program.",
            "Residence or household evidence only when location or financial need is mandatory.",
            "Guardian relationship or authorization evidence when required for final validation, preferably deferred until the learner is shortlisted when risk allows.",
        ],
        [
            ("Guardian reviews the program. ", "Confirm the benefit, school coverage, age/grade rule, child-safeguarding arrangements, privacy notice, and any activity expectations."),
            ("Create a guardian-supported learner record. ", "Separate the child's information from the adult's login and contact record."),
            ("Complete a simple eligibility check. ", "Use age, grade, school, residence, and need criteria without asking the child to interpret policy language."),
            ("Upload minimum evidence. ", "Provide school and need evidence required for pre-screening; keep highly sensitive originals offline unless later requested through a secure process."),
            ("Guardian submits and receives updates. ", "The platform sends understandable notices to the authorized adult while keeping a record linked to the learner."),
            ("Provider reviews with child-sensitive safeguards. ", "Interviews or assessments should be age-appropriate, supervised, scheduled safely, and limited to relevant questions."),
            ("Guardian accepts the offer. ", "The provider explains benefit delivery, school coordination, attendance, renewal, and any photo or publicity consent separately."),
        ],
        [
            "Use the minimum data needed and avoid public display of the child's identity, address, school schedule, or contact details.",
            "Do not bundle scholarship consent with optional publicity, marketing, photograph, or testimonial consent.",
            "Provide assisted and offline channels for families with low connectivity or digital literacy.",
            "A child should never be asked for a password, one-time PIN, bank credential, or unsupervised meeting with an unknown reviewer.",
        ],
        [
            "The adult cannot establish authority to act for the learner when the program requires it.",
            "The learner falls outside the stated age, grade, school, or location coverage.",
            "Required school or household evidence is missing or inconsistent and is not corrected by the deadline.",
            "The family cannot participate in a mandatory program component and no reasonable alternative is available.",
        ],
    )

    add_applicant_process(
        document,
        "18. Junior high school learners",
        "The learner can participate in completing the application, but a parent or guardian normally supports consent, document handling, schedules, and important decisions while the learner is a minor.",
        "The process begins to use academic records and learner interests more directly, but it must remain suitable for minors and should not demand college-level documents or career certainty.",
        [
            "Current grade, school, school type, location, academic period, and learner interests.",
            "Guardian contact and relationship, household context when relevant, and accessibility needs.",
            "Any target participation such as STEM enrichment, arts, sports, leadership, or transition support.",
        ],
        [
            "Latest report card or school-certified grades when academic standing is a criterion.",
            "Enrollment or school certification and location/income proof only when required.",
            "Portfolio, recommendation, or activity record only for a talent, leadership, or service scholarship.",
        ],
        [
            ("Review the target group. ", "Confirm eligible grade levels, school coverage, location, and whether the program continues into senior high school."),
            ("Complete the learner profile with guardian support. ", "Record current education and contact information using consistent names."),
            ("Check academic and non-academic rules. ", "The system should explain whether a rule is mandatory or used only for ranking."),
            ("Prepare and submit evidence. ", "Upload the latest relevant school record rather than unrelated sensitive documents."),
            ("Join an age-appropriate assessment if announced. ", "The guardian receives schedule and safety information; accommodations are requested early."),
            ("Receive the decision and transition plan. ", "If selected, understand benefit release, school coordination, continuing grade expectations, and the next school-year process."),
        ],
        [
            "Keep guardian and learner communication roles clear; do not expose the learner's contact details to other applicants.",
            "Allow equivalent achievements for learners from schools with fewer clubs or competitions.",
            "Avoid penalizing a learner for a grading scale difference that the provider has not normalized fairly.",
        ],
        [
            "The submitted academic period is not the one required by the program.",
            "The learner is outside eligible grade levels or participating schools.",
            "A required guardian action, correction, exam, or interview is missed.",
            "Portfolio or activity claims cannot be supported when they are material to selection.",
        ],
    )

    add_applicant_process(
        document,
        "19. Senior high school learners",
        "The learner normally completes much of the application, with parent or guardian involvement when the learner is a minor and when consent, financial evidence, or final commitments require it.",
        "Track, strand, current grade, transition plans, and school participation may affect fit. A Senior High School voucher is a subsidy arrangement and should not automatically be described as a private scholarship or cash award.",
        [
            "Grade 11 or 12 status, school, track/strand, subjects or academic average, and intended next pathway.",
            "Household and location information for need- or coverage-based programs.",
            "Availability for enrichment, review, work immersion, exam, interview, or mentoring when announced.",
        ],
        [
            "Report card, certificate of grades, enrollment/admission record, or school certification.",
            "Income and residence evidence when mandatory.",
            "Portfolio, competition result, recommendation, or project evidence for specialized programs.",
        ],
        [
            ("Identify the support category. ", "Determine whether the opportunity is a voucher, tuition discount, scholarship, allowance, training, or combined package."),
            ("Check track/strand and grade coverage. ", "Equivalent strand or course labels should be recognized when the provider permits them."),
            ("Complete profile and upload current records. ", "Use the grading scale shown by the school and identify the covered period."),
            ("Submit before the published deadline. ", "Review all commitments, including attendance, mentoring, research, or transition activities."),
            ("Complete competitive steps. ", "Take any announced exam, interview, portfolio review, or orientation through official channels."),
            ("Accept and coordinate with the school. ", "Confirm whether the provider pays the school, reimburses expenses, or releases benefits directly."),
            ("Renew or transition. ", "Provide current grades and enrollment if renewable, and understand whether support continues into TVET or college."),
        ],
        [
            "Explain that strand matching is a fit indicator, not proof of final selection.",
            "For minors, route major decisions and sensitive requests through the authorized guardian without excluding the learner from understandable communication.",
            "Publish whether changing school, track, or strand affects the award.",
        ],
        [
            "Track/strand, grade, school, or location falls outside a mandatory rule.",
            "The application confuses voucher eligibility with the scholarship's separate requirements.",
            "The learner misses an assessment or does not respond to a correction request.",
            "The selected learner changes pathway without obtaining approval required by the program.",
        ],
    )

    add_applicant_process(
        document,
        "20. Alternative Learning System learners and completers",
        "The learner normally leads the application, supported by an ALS teacher, learning facilitator, guardian, or authorized community contact when appropriate.",
        "ALS pathways may not use the same year-level, semester, or transcript structure as formal schooling. The provider should accept valid ALS participation, portfolio, assessment, or equivalency evidence instead of forcing conventional school fields.",
        [
            "ALS program, learning center, current participation or completion status, equivalency level, and intended next pathway.",
            "Prior education history only to the extent needed to establish eligibility.",
            "Work, family, connectivity, transport, and schedule constraints relevant to participation.",
        ],
        [
            "ALS enrollment or participation certification, portfolio evidence, assessment result, or equivalency credential as applicable.",
            "Admission or training record for the next intended pathway when required.",
            "Income/residence evidence for targeted assistance.",
        ],
        [
            ("Confirm that ALS applicants are included. ", "Do not assume a program for 'students' automatically accepts every equivalency pathway."),
            ("Use an ALS-compatible profile. ", "Record equivalency level and learning center rather than requiring an inaccurate conventional grade."),
            ("Map the intended transition. ", "Identify whether support is for ALS completion, SHS, TVET, college, employment preparation, or another next step."),
            ("Submit relevant evidence. ", "Use the official or facilitator-certified record named by the provider."),
            ("Complete accessible review steps. ", "Offer reasonable scheduling and assisted submission where work or connectivity creates barriers."),
            ("Validate the next-pathway condition. ", "Before release, confirm admission, enrollment, or training placement if the award depends on it."),
        ],
        [
            "Provide 'ALS' and equivalent statuses in forms, filters, and matching logic.",
            "Do not mark an applicant mismatched only because a GWA, strand, or conventional school type is not applicable.",
            "Allow alternative proof when a formal record is legitimately unavailable and document how it was verified.",
        ],
        [
            "The scholarship excludes the relevant ALS or transition level.",
            "The submitted credential does not establish the equivalency required for the next pathway.",
            "The learner cannot meet a mandatory enrollment or training start date.",
        ],
    )

    add_applicant_process(
        document,
        "21. TVET applicants",
        "The applicant normally leads, while the chosen Technical Vocational Institution and program administrator may validate admission, qualification, training, assessment, and attendance.",
        "TVET support is tied to a training regulation, qualification, approved training provider, competency assessment, certification, and often employment outcomes. Conventional college course and GWA fields may be irrelevant.",
        [
            "Target qualification, competency level, training provider, delivery mode, schedule, location, and prior competencies.",
            "Employment status, income or priority sector when relevant, and availability for training and assessment.",
            "Existing National Certificates, Certificates of Competency, or recognition of prior learning where applicable.",
        ],
        [
            "Admission or enrollment record from an eligible training provider.",
            "Identity, education or prerequisite evidence required by the qualification.",
            "Priority-sector, income, employment, or residence evidence when specified.",
            "Assessment, certification, attendance, or completion records for release and closure.",
        ],
        [
            ("Choose the qualification and provider. ", "Confirm that the training program and institution are eligible under the funding call."),
            ("Check prerequisites and schedule. ", "Review minimum education, age if lawful and relevant, tools, health/safety requirements, and time commitment."),
            ("Apply through the designated channel. ", "Some programs use TESDA or provider systems in addition to a finder platform."),
            ("Complete validation and orientation. ", "The provider confirms slot availability, required records, training rules, and assessment obligations."),
            ("Attend training and competency assessment. ", "Attendance and competency evidence may control payment or continued support."),
            ("Receive certification and transition support. ", "The program records completion and may provide employment referral, entrepreneurship support, or further training."),
        ],
        [
            "Show training and assessment costs separately and state whether tools, allowance, and certification are covered.",
            "Do not promise employment unless placement is genuinely guaranteed; describe referrals and partnerships accurately.",
            "Support working learners with clear schedules and attendance consequences.",
        ],
        [
            "The selected qualification or training provider is not covered.",
            "Prerequisites, schedule, or assessment conditions cannot be met.",
            "The applicant is already funded for the same training in a prohibited combination.",
            "Attendance or competency requirements for continued support are not achieved.",
        ],
    )

    add_applicant_process(
        document,
        "22. Incoming undergraduate applicants",
        "The applicant normally leads. A parent or guardian may support financial documents, but the applicant should understand and control the application if legally capable.",
        "The learner may be applying before final enrollment. Programs must distinguish admission, intention to enroll, and actual registration, and should state whether an offer is conditional on entering an eligible institution or course.",
        [
            "Senior high school record, graduating status, intended degree, preferred institution, admission status, and entrance assessment information where relevant.",
            "Household income, residence, and existing financial aid when required.",
            "Priority field, talent, leadership, service, disability accommodation, or other criteria used by the program.",
        ],
        [
            "Latest report card, certificate of grades, or graduation/completion evidence.",
            "Admission notice, application receipt, or later certificate of enrollment as defined by the program.",
            "Income and residence evidence for need/coverage rules.",
            "Program-specific portfolio, recommendation, exam result, or declaration of existing aid.",
        ],
        [
            ("Check admission timing. ", "Determine whether the scholarship application comes before or after school admission and which institutions/courses qualify."),
            ("Complete the profile and match review. ", "Record intended course names using standardized and equivalent labels where possible."),
            ("Submit preliminary evidence. ", "Provide the latest available academic record and indicate which final documents are still pending."),
            ("Complete competitive review. ", "Take any exam/interview and respond to evidence corrections."),
            ("Receive a conditional or final outcome. ", "A conditional offer should list the exact admission/enrollment evidence and deadline needed to confirm it."),
            ("Enroll and complete onboarding. ", "Submit final registration through a secure channel, accept terms, and confirm payment or allowance arrangements."),
        ],
        [
            "Do not label an applicant ineligible merely because a final enrollment document cannot exist before school registration, unless current enrollment is an explicit rule.",
            "Explain course equivalencies and how a change of institution or course affects the offer.",
            "Identify funding gaps where the scholarship is partial.",
        ],
        [
            "The applicant is not admitted to or enrolled in an eligible institution/course by the confirmation deadline.",
            "Final grades or graduation status do not meet a conditional requirement.",
            "Another award creates a prohibited overlap and is not resolved.",
            "The applicant misses acceptance or enrollment confirmation.",
        ],
    )

    add_applicant_process(
        document,
        "23. Current undergraduate, transferee, and returning applicants",
        "The applicant normally leads, and the current or receiving institution may validate enrollment, grades, units, academic standing, and transfer status.",
        "Current students have a record of tertiary performance and remaining study time. Transferees and returning learners may have interrupted, shifted, or multi-institution records that require context rather than an automatic mismatch.",
        [
            "Current institution, degree, year level, units, GWA/GPA and grading scale, academic standing, expected completion, and any course/school change.",
            "Financial and location information when required.",
            "Existing scholarships, leave history, employment, caregiving, or other circumstances relevant to continuation.",
        ],
        [
            "Registration form or certificate of enrollment, official or certified grades, curriculum/remaining units where required.",
            "Transfer credentials, acceptance notice, or readmission evidence for transferees/returnees.",
            "Income, residence, and other program-specific evidence.",
        ],
        [
            ("Confirm current-student eligibility. ", "Check year level, remaining terms, course/institution coverage, minimum units, and academic standing."),
            ("Explain unusual records. ", "Declare approved leave, transfer, shift, incomplete marks, or grading-scale differences accurately."),
            ("Submit the current academic period. ", "Use the exact period and scale required by the program."),
            ("Complete ranking or interview. ", "Providers may consider persistence, completion plan, need, and program fit in addition to grades."),
            ("Accept and coordinate release. ", "Confirm whether prior expenses are reimbursable or only future approved costs are covered."),
            ("Meet renewal conditions. ", "Submit enrollment and performance evidence each stated period and request approval before material changes."),
        ],
        [
            "Allow explanations and reviewer notes for shifts, transfers, non-standard terms, and legitimate interruptions.",
            "A system should not average unlike grading scales without a documented method.",
            "Publish rules for leave, reduced load, internship terms, practicum, and final-year status.",
        ],
        [
            "Insufficient remaining study period or units under the program rules.",
            "Academic standing, course, institution, or enrollment load is outside a mandatory threshold.",
            "Records from prior institutions are incomplete and material to review.",
            "An unapproved course/school change affects eligibility.",
        ],
    )

    add_applicant_process(
        document,
        "24. Graduate, postgraduate, and research applicants",
        "The applicant leads and may coordinate with a graduate school, employer, faculty supervisor, research adviser, ethics body, or host institution.",
        "Advanced-study support often evaluates a study or research plan, institutional capacity, public or professional relevance, outputs, and milestones rather than relying mainly on household need or an undergraduate-style grade cutoff.",
        [
            "Degree level, field, institution, admission status, employment/faculty status, proposed study mode, and timeline.",
            "Research title, problem, methods, adviser, host resources, ethics needs, outputs, and budget where applicable.",
            "Prior degree record, professional experience, publications, service plan, and other funding.",
        ],
        [
            "Admission or nomination, transcript, curriculum vitae, recommendations, and employer endorsement when required.",
            "Research proposal, work plan, budget, ethics clearance or plan, and adviser/host confirmation.",
            "Language, passport/visa, or international-host evidence for mobility programs.",
        ],
        [
            ("Confirm the award model. ", "Identify whether support is a scholarship, fellowship, research grant, faculty development award, or mobility program."),
            ("Establish institutional and adviser fit. ", "Confirm admission, supervision, facilities, and any host agreement."),
            ("Submit a complete technical and personal package. ", "Separate academic/research scoring from administrative eligibility."),
            ("Undergo expert review and interview. ", "Use reviewers with relevant competence and manage conflicts of interest."),
            ("Negotiate and accept conditions. ", "Clarify intellectual property, publication, data, ethics, travel, return service, and changes to scope."),
            ("Report milestones and expenditures. ", "Release later tranches only against the announced academic, technical, or financial requirements."),
            ("Close outputs and obligations. ", "Submit degree, publication, research, service, or knowledge-transfer evidence."),
        ],
        [
            "Protect unpublished proposals and sensitive research data; limit access to authorized reviewers.",
            "Disclose intellectual-property and publication conditions before acceptance.",
            "Allow formal change requests for adviser, method, timeline, institution, or research scope.",
        ],
        [
            "Admission, adviser, host, or research feasibility is not established.",
            "The proposal falls outside priority areas or fails ethics/safety requirements.",
            "The budget is unsupported or duplicates another funding source.",
            "Required milestones or approved changes are not completed.",
        ],
    )

    add_applicant_process(
        document,
        "25. Working learners, returning learners, and out-of-school youth",
        "The applicant leads, with possible support from an employer, ALS facilitator, school, training provider, social worker, or community organization.",
        "These applicants may have interrupted records, limited time, informal employment, caregiving responsibilities, or no current enrollment. A program should define whether it supports re-entry, current study, certification, or transition rather than assuming continuous full-time schooling.",
        [
            "Last completed level, interruption reason if relevant and voluntarily provided, target pathway, admission/enrollment plan, and expected schedule.",
            "Work and caregiving constraints, income pattern, transport/connectivity, and prior competencies when relevant.",
            "Support needed to re-enter, remain, or complete.",
        ],
        [
            "Prior school/ALS/training records, admission or re-entry record, employment or income evidence when required.",
            "Recognition of prior learning, portfolio, or competency evidence where the pathway permits it.",
            "Alternative income evidence for informal or irregular work when accepted.",
        ],
        [
            ("Choose a program that accepts the actual status. ", "Confirm whether current enrollment is required at application or only before release."),
            ("Reconstruct the education pathway. ", "Use the latest valid records and explain gaps without demanding irrelevant personal details."),
            ("Assess practical feasibility. ", "Review schedule, transport, care, connectivity, and uncovered costs."),
            ("Submit and complete assisted review. ", "Use help channels when records or digital access are limited."),
            ("Confirm re-entry or training placement. ", "A conditional award may become final after enrollment or admission."),
            ("Monitor continuation. ", "Use proportionate milestones suitable for part-time, modular, ALS, TVET, or returning pathways."),
        ],
        [
            "Do not treat an education gap itself as misconduct or lack of merit.",
            "Offer alternate proofs for informal work or lost records when policy permits.",
            "State attendance and scheduling requirements before submission.",
        ],
        [
            "The program requires current enrollment and does not allow a conditional re-entry pathway.",
            "The chosen schedule or institution is outside program coverage.",
            "Prior records cannot establish a required prerequisite and no alternative assessment exists.",
        ],
    )

    add_applicant_process(
        document,
        "26. Applicants with disabilities or support needs",
        "The applicant leads when capable; a parent, guardian, or authorized support person may assist. Support should increase access without removing the applicant's autonomy.",
        "Disability may be an eligibility category, an accommodation need, neither, or both. Providers should not collect diagnosis details unless directly needed for eligibility, accommodation, safety, or a lawful program purpose.",
        [
            "Education pathway and ordinary eligibility information.",
            "Functional accommodation needs for forms, documents, assessments, travel, communication, or program participation.",
            "Disability category or proof only when the scholarship specifically targets it.",
        ],
        [
            "The same academic, enrollment, income, or location evidence required from comparable applicants.",
            "A narrowly scoped disability or accommodation record only when necessary.",
            "Alternative-format records or authorized support-person information when requested by the applicant.",
        ],
        [
            ("Review accessibility before applying. ", "Check form compatibility, venue access, communication methods, schedules, and benefit coverage."),
            ("Request reasonable accommodation. ", "State the functional need and preferred adjustment without unnecessary medical detail."),
            ("Complete the standard eligibility process. ", "Disability should not create unrelated additional barriers."),
            ("Arrange accessible assessment. ", "Provide reasonable time, format, assistive technology, interpreter, reader, or venue adjustments as appropriate."),
            ("Confirm ongoing support. ", "Clarify who supplies accommodations and whether related costs are covered."),
        ],
        [
            "Use accessible interfaces, keyboard navigation, readable contrast, alt text, plain language, and non-digital assistance.",
            "Keep disability records separate from general browsing and restrict reviewer access.",
            "Assess the published criteria, not assumptions about capacity.",
        ],
        [
            "A genuinely essential program condition cannot be met even with a reasonable adjustment, and the provider documents the reason.",
            "Required evidence for a targeted disability award is missing and not corrected.",
            "The applicant cannot access the process because no accommodation route was offered; this is a process problem that should be escalated, not silently treated as applicant failure.",
        ],
    )

    add_applicant_process(
        document,
        "27. Indigenous, geographically isolated, displaced, or disaster-affected applicants",
        "The applicant leads when possible, with culturally appropriate support from a guardian, school, community representative, LGU, social worker, or authorized partner.",
        "Distance, displacement, damaged records, language, unstable connectivity, and community data sensitivity may make a normal online process unfair or impossible. Programs should define acceptable alternate evidence and assisted channels in advance.",
        [
            "Current location and usual residence, education pathway, displacement or access condition only when relevant, and safe contact method.",
            "School or learning-center connection, target institution, travel/boarding needs, and language/access needs.",
            "Community affiliation only when a stated program criterion requires it.",
        ],
        [
            "School, LGU, social welfare, community, or displacement certification appropriate to the program.",
            "Alternative academic or identity records where originals were lost, subject to later validation.",
            "Residence/community evidence that does not expose sensitive location data unnecessarily.",
        ],
        [
            ("Use an accessible discovery channel. ", "Provide official mobile, assisted, school, LGU, or community access rather than relying only on high-bandwidth web forms."),
            ("Confirm flexible evidence rules. ", "Identify accepted alternatives before submission."),
            ("Complete a minimum-data pre-screen. ", "Request only information needed to determine program fit and immediate support."),
            ("Validate through authorized partners. ", "Use a documented verification path while avoiding public disclosure of sensitive community records."),
            ("Plan delivery. ", "Account for travel, boarding, connectivity, school access, and safe benefit distribution."),
            ("Monitor without creating impossible reporting. ", "Use proportional evidence and alternate communication during disruption."),
        ],
        [
            "Do not publish exact household or community coordinates.",
            "Support local language or plain-language explanation where feasible.",
            "Do not reject solely because the standard document is unavailable when the policy permits a credible alternative.",
        ],
        [
            "The coverage area is mandatory and the applicant is outside it.",
            "Alternative evidence cannot establish a material criterion after fair attempts to verify.",
            "The program cannot safely or lawfully deliver the stated benefit to the location and did not disclose this limitation.",
        ],
    )

    add_applicant_process(
        document,
        "28. International study and mobility applicants",
        "The applicant leads and coordinates with the home institution, host institution, scholarship provider, immigration authorities, and guardian when the applicant is a minor.",
        "These programs add cross-border admission, passport/visa, travel, language, health, insurance, financial proof, recognition, and data-transfer requirements. Selection is not the same as permission to enter or remain in the destination country.",
        [
            "Home and host institution, program, dates, language, nationality/citizenship, travel readiness, and emergency contact.",
            "Academic record, study/research plan, funding gap, other grants, and return plan.",
            "Accommodation, disability, health, safety, and dependent considerations only as necessary.",
        ],
        [
            "Passport, admission/invitation, transcript, language evidence, recommendation, and study or research plan.",
            "Visa, insurance, medical or vaccination evidence, financial proof, and travel documents at the correct stage.",
            "Parental authorization and safeguarding arrangements for a minor.",
        ],
        [
            ("Verify the provider and host. ", "Confirm official domains, accreditation/recognition, funding terms, and scam warnings."),
            ("Separate award selection from mobility clearance. ", "Identify admission, visa, health, insurance, and travel conditions."),
            ("Submit academic and mobility evidence in stages. ", "Avoid collecting passport and medical data before needed."),
            ("Complete interview and nomination. ", "Understand whether the provider nominates candidates or makes the final award."),
            ("Secure admission, visa, and insurance. ", "Meet external authority deadlines that the scholarship cannot waive."),
            ("Travel, report, and return. ", "Follow safety, academic, financial, and return obligations."),
        ],
        [
            "Use secure channels for passport, visa, medical, and financial data.",
            "State uncovered airfare, deposits, exchange-rate risk, dependent costs, and emergency responsibility.",
            "Explain cross-border data recipients and retention where applicable.",
        ],
        [
            "The applicant receives a scholarship nomination but not host admission or a visa.",
            "Language, travel, health, insurance, or financial conditions are not completed.",
            "Costs outside the award make participation infeasible and no supplementary funding is available.",
        ],
    )

    document.add_heading("29. Summary matrix by applicant type", level=2)
    add_table(
        document,
        ["Applicant type", "Primary lead", "Core evidence", "Main process adjustment"],
        [
            ("Kindergarten and elementary", "Guardian", "Age/grade, school, relevant need", "Child-safe, minimum-data, guardian-led process"),
            ("Junior high", "Learner + guardian", "Current school record and targeted evidence", "Age-appropriate assessment and communication"),
            ("Senior high", "Learner + guardian if minor", "Grade/strand, enrollment, relevant need", "Distinguish voucher, scholarship, and transition support"),
            ("ALS", "Learner", "ALS participation/equivalency and next pathway", "Accept non-conventional education records"),
            ("TVET", "Learner + training provider", "Qualification admission and prerequisites", "Focus on training, assessment, and certification"),
            ("Incoming undergraduate", "Learner", "SHS record and admission/enrollment when available", "Allow conditional award pending final enrollment"),
            ("Current undergraduate", "Learner + institution", "Enrollment, grades, units, academic standing", "Handle transfer, shift, leave, and renewal"),
            ("Graduate/research", "Applicant + host/adviser", "Admission, proposal, references, milestones", "Expert review, ethics, outputs, and service"),
            ("Working, returning, or OSY", "Learner", "Prior record and re-entry plan", "Flexible records, schedules, and conditional enrollment"),
            ("Disability or support needs", "Applicant with chosen support", "Ordinary evidence plus necessary accommodation proof", "Accessible process and data minimization"),
            ("Isolated, displaced, or disaster-affected", "Applicant + authorized local support", "Alternative school/residence/status evidence", "Assisted channels and alternate verification"),
            ("International or mobility", "Applicant + institutions", "Admission, passport/visa at proper stage", "Separate scholarship decision from immigration clearance"),
        ],
        widths=[Inches(1.35), Inches(1.25), Inches(2.05), Inches(2.1)],
        font_size=8.5,
    )

    document.add_heading("Part VI. Selection, Provider Operations, and Decisions", level=1)

    document.add_heading("30. Selection methods", level=2)
    add_bullets(document, [
        ("Pass/fail eligibility review. ", "Each mandatory condition is met, not met, unknown, or not applicable. 'Any' and blank unrestricted criteria must not create a mismatch."),
        ("Weighted rubric. ", "Relevant dimensions receive published or internally approved weights. Reviewers enter evidence-based scores and comments rather than choosing an unexplained overall status."),
        ("Ranking. ", "Eligible applicants are ordered using a consistent score and tie-break rules when qualified demand exceeds slots."),
        ("Exam. ", "A test measures specified knowledge, aptitude, or readiness. Format, scope, accessibility, identity controls, passing rule, and result use should be stated."),
        ("Interview. ", "Structured questions and anchored scoring improve consistency. Reviewers should avoid irrelevant personal or discriminatory questions."),
        ("Portfolio, audition, or demonstration. ", "Evidence is assessed against a rubric appropriate to talent, practice, research, or innovation."),
        ("Lottery among eligible applicants. ", "A transparent random method may be appropriate when the purpose does not justify ranking and eligible demand exceeds slots."),
        ("Hybrid process. ", "Many programs combine minimum eligibility, document verification, ranking, and an exam or interview."),
    ])

    document.add_heading("A defensible rubric", level=3)
    add_bullets(document, [
        "Uses criteria connected to the stated program goal and excludes information not needed for the decision.",
        "Defines score anchors so reviewers know what weak, adequate, strong, and exceptional evidence means.",
        "Requires a score and reason for each criterion rather than an unexplained total.",
        "Records reviewer identity, time, conflicts, overrides, and the evidence considered.",
        "Uses separate eligibility and competitive scores so a high merit score cannot silently override a mandatory rule.",
        "Defines tie-breakers, minimum passing conditions, waitlist order, and authority for final approval.",
        "Is tested on sample cases for unintended disadvantage, inconsistent interpretation, and data errors.",
    ])

    document.add_heading("31. Provider responsibilities from design to closure", level=2)
    add_numbered(document, [
        ("Authorize the program. ", "Confirm the legal organization, responsible officer, funding source, budget, reviewer authority, and contact channel."),
        ("Design a coherent offer. ", "Align goal, target group, benefit, criteria, documents, rubric, schedule, and obligations."),
        ("Publish complete information. ", "Avoid vague benefits, undisclosed screening, hidden documents, and misleading guarantees."),
        ("Protect applicant data. ", "Use role-based access, least privilege, secure file routes, retention schedules, incident procedures, and staff confidentiality."),
        ("Review consistently. ", "Train reviewers, disclose conflicts, use the approved rubric, and permit specific corrections."),
        ("Communicate actionable notices. ", "Every notice should say what happened, what the applicant must do, the deadline, and where to ask for help."),
        ("Complete final due diligence. ", "Validate originals or official records only when necessary and through an authorized method."),
        ("Release benefits accountably. ", "Document direct payments, transfers, goods, services, acknowledgements, failed releases, and corrections."),
        ("Monitor and resolve cases. ", "Apply renewal, suspension, withdrawal, grievance, waitlist, and replacement-recipient rules fairly."),
        ("Evaluate and close. ", "Measure reach, completion, service quality, and program outcomes using minimized and properly retained data."),
    ])

    document.add_heading("32. Exam, interview, formal application, and award order", level=2)
    document.add_paragraph(
        "There is no universal order. A provider may require a formal provider application before an exam, after a pre-screen, or only from finalists. What matters is that the order is disclosed, each stage has a clear purpose, and applicants do not repeatedly submit the same information without need."
    )
    add_table(
        document,
        ["Model", "Suitable when", "Recommended flow"],
        [
            ("Document-first", "Mandatory eligibility can be established cheaply from existing records.", "Pre-screen documents -> eligibility review -> exam/interview -> final validation -> decision"),
            ("Assessment-first", "An accessible low-cost assessment is the main gateway and sensitive documents should be deferred.", "Basic profile -> assessment -> shortlist -> documents -> interview/formal application -> decision"),
            ("Formal-application handoff", "The finder is intentionally only a prescreen and the provider has a separate legal or institutional application.", "Platform prescreen -> qualified notice -> provider form/originals -> provider stages -> final decision"),
            ("Portfolio/research", "Quality and fit cannot be judged from grades alone.", "Eligibility -> proposal/portfolio -> expert review -> interview -> due diligence -> award"),
            ("Emergency assistance", "Delay could cause immediate dropout or harm.", "Minimum eligibility -> rapid verification -> provisional decision -> follow-up evidence -> closure"),
        ],
        widths=[Inches(1.4), Inches(2.55), Inches(2.8)],
    )

    document.add_heading("33. Offers, waitlists, alternates, and unsuccessful outcomes", level=2)
    add_bullets(document, [
        ("Offer. ", "State the full support package, duration, conditions, start date, acceptance deadline, release route, and contact. Avoid describing a conditional offer as money already received."),
        ("Waitlist. ", "State the order or method, validity period, whether further action is needed, and that a place is not guaranteed."),
        ("Alternate recipient. ", "Use the pre-approved waitlist or ranking when a selected applicant declines or fails final conditions. Record the reason and authority."),
        ("Not selected. ", "Use respectful language. Where policy allows, state whether the result came from ineligibility, incomplete evidence, assessment result, or limited ranking."),
        ("Appeal or reconsideration. ", "Limit appeals to stated grounds such as process error, overlooked evidence, or factual correction. An appeal is not a new application unless policy says so."),
    ])

    document.add_heading("34. Renewal, changes, suspension, and termination", level=2)
    add_bullets(document, [
        "Renewal criteria may include current enrollment, minimum academic or competency standing, required units, attendance, outputs, financial need, conduct, and available funding.",
        "The program should define how incomplete grades, illness, disability, disaster, leave, reduced load, internship, practicum, transfer, or course change are handled.",
        "A change that affects eligibility should be declared promptly and reviewed before the next release when possible.",
        "Suspension temporarily pauses benefits while a condition is resolved; termination ends the award. The provider should state grounds, notice, effect on unpaid benefits, and any review route.",
        "Recovery of funds should occur only under stated lawful conditions and should distinguish fraud from an honest error or provider processing mistake.",
        "Recipients should receive a written status and should not discover suspension only when an expected payment fails.",
    ])

    document.add_heading("Part VII. Privacy, Safety, Fairness, and Problem Handling", level=1)

    document.add_heading("35. Privacy principles for scholarship data", level=2)
    document.add_paragraph(
        "Scholarship applications can contain education records, household finances, addresses, identity documents, disability information, and data about minors. The Philippine Data Privacy Act framework emphasizes transparency, legitimate purpose, and proportionality. In practical terms, the provider and platform should explain what is collected, why it is needed, who can see it, how long it is kept, and how the applicant can exercise applicable rights."
    )
    add_bullets(document, [
        ("Transparency. ", "Give a readable privacy notice before collection and a more detailed notice that identifies the organization, purposes, recipients, retention, security contact, and rights."),
        ("Legitimate purpose. ", "Use data for defined scholarship, verification, communication, release, monitoring, safeguarding, and accountability purposes rather than unrelated marketing."),
        ("Proportionality. ", "Collect information that is adequate and relevant but not excessive for the present stage."),
        ("Role-based access. ", "Applicants see their own records; provider reviewers see only assigned or authorized cases; administrators oversee legitimate platform functions without casual access."),
        ("Retention and disposal. ", "Keep unsuccessful and recipient records only for documented operational, legal, audit, or dispute periods, then securely delete or anonymize them."),
        ("Accuracy and correction. ", "Let applicants correct factual errors while preserving submitted versions and review history where accountability requires it."),
        ("Incident response. ", "Record suspected unauthorized access, contain it, assess risk, reset credentials where needed, and make required notifications."),
    ])

    document.add_heading("Minors and guardian-supported applications", level=3)
    add_bullets(document, [
        "Separate the learner's profile from the adult representative's account and record the relationship and authority.",
        "Use an age-appropriate notice for the learner and a complete notice for the adult making or supporting the application.",
        "Do not assume that a guardian may consent to unrelated marketing, publicity, or data sharing merely because the guardian supports the scholarship application.",
        "Keep learner records confidential and share them only with authorized staff, providers, schools, or partners for a disclosed purpose.",
        "When legal capacity, custody, or authority is uncertain, pause the sensitive action and use an approved verification or safeguarding process.",
    ])

    document.add_heading("36. Security and scam prevention", level=2)
    add_bullets(document, [
        "Confirm the provider through an official website, verified portal profile, government or institutional directory, or independently obtained contact.",
        "Be cautious when an award is guaranteed, the deadline is artificially urgent, the provider refuses written terms, or the benefit is unrealistically large.",
        "Do not send passwords, one-time PINs, full card credentials, or remote-control access. A legitimate administrator does not need the applicant's password.",
        "Treat requests for upfront 'processing,' 'slot reservation,' or release fees as high risk unless the official program clearly identifies a lawful fee and recipient. Scholarships ordinarily should not require payment simply to win an award.",
        "Use protected upload routes rather than public file links or social-media messages for sensitive evidence.",
        "Check email sender domains and destination links; a display name alone does not prove authenticity.",
        "Report suspicious programs to the platform and appropriate provider or authority without confronting an unknown actor using personal information.",
    ])

    document.add_heading("37. Fairness and decision-support safeguards", level=2)
    document.add_paragraph(
        "A decision support system should make criteria easier to apply, not hide the decision. Its score is guidance for discovery and review. The provider remains accountable for eligibility rules, rubric weights, evidence, overrides, and final outcomes."
    )
    add_bullets(document, [
        "Show why a criterion is met, not met, unknown, or not applicable in plain language.",
        "Treat unrestricted values such as 'Any education level,' 'Any course,' 'Any location,' or 'No academic minimum' as neutral matches, never as failures.",
        "Normalize synonymous courses, strands, municipalities, school types, and grading formats through maintained reference lists plus a reviewable custom option.",
        "Do not infer sensitive attributes from names, photographs, addresses, or documents when they are not explicit program criteria.",
        "Require human review before a negative decision based on extracted text, distance, identity matching, or incomplete profile data.",
        "Record model/rule version, input values, explanation, reviewer action, and override reason so errors can be investigated.",
        "Test for false mismatches across education levels, ALS/TVET paths, disabilities, grading systems, equivalent course names, and applicants with missing but non-mandatory data.",
    ])

    document.add_heading("38. Complaints, reports, and conflicts", level=2)
    add_table(
        document,
        ["Concern", "Primary handler", "Resolution approach"],
        [
            ("Program information or provider conduct", "Provider and platform administrator", "Preserve the report; provider answers program facts while the administrator handles platform safety and escalation."),
            ("Application review or correction", "Provider", "Use the program's review record, criteria, evidence, deadline, and reconsideration policy."),
            ("Platform access or technical failure", "Platform administrator", "Restore access, preserve submission timing, and avoid penalizing applicants for confirmed system failure."),
            ("Privacy request or suspected exposure", "Responsible privacy/security contact", "Verify identity proportionately, restrict further access, investigate, document, and follow applicable notice duties."),
            ("Payment or benefit release", "Provider or paying institution", "Trace the approved benefit, release record, destination, failed transaction, and correction authority."),
            ("Fraud or falsification", "Authorized provider/admin review", "Restrict access only as necessary, preserve evidence, give fair notice, and escalate under written policy."),
        ],
        widths=[Inches(1.6), Inches(1.65), Inches(3.5)],
    )
    document.add_paragraph(
        "When both a provider and administrator can resolve a report, the system should use one owner, one shared status history, and role-specific actions. A provider response should not erase an administrator's safety review, and an administrator closure should record whether provider action remains outstanding."
    )

    document.add_heading("Part VIII. Practical Checklists, Glossary, and Official References", level=1)

    document.add_heading("39. Applicant checklist", level=2)
    add_bullets(document, [
        "I confirmed that the provider and application channel are authentic.",
        "I understand the exact benefit, what is not covered, duration, release method, and renewal conditions.",
        "I meet every mandatory criterion or have identified an item that requires provider clarification.",
        "My profile uses the correct education pathway, grading scale, course/strand, location, and contact details.",
        "My required files are readable, current, relevant, and consistent; I have kept the originals secure.",
        "I understand which documents are required now and which may be requested only at formal application or award stage.",
        "I disclosed other assistance when required and checked compatibility rules.",
        "I understand the selection stages, dates, venue or online mode, accommodation route, and final decision authority.",
        "I read the privacy notice, application terms, correction and withdrawal rules, and complaint channel.",
        "I saved my application reference and will rely on official notifications rather than unofficial promises.",
    ])

    document.add_heading("40. Parent or guardian checklist", level=2)
    add_bullets(document, [
        "I am authorized to create or support this application and the learner record is separate from my representative account.",
        "The learner received an age-appropriate explanation and understands the main purpose and expected activities.",
        "Only information necessary for the scholarship was provided, and optional publicity or marketing permission was not treated as mandatory.",
        "I verified all meeting, assessment, transport, venue, safeguarding, and contact arrangements.",
        "I understand how the benefit reaches the learner or school and what expenses remain our responsibility.",
        "I know how to correct information, withdraw, report a concern, or request an accommodation.",
    ])

    document.add_heading("41. Provider program-quality checklist", level=2)
    add_bullets(document, [
        "The purpose, target group, criteria, evidence, rubric, benefits, slots, budget, and process all support the same objective.",
        "Every field and document has a defined current-stage purpose; sensitive originals are deferred when possible.",
        "Different education pathways and grading systems are handled accurately, with 'Any' options working as unrestricted.",
        "Applicants can see selection stages, schedules, formal handoff, corrections, withdrawal, waitlist, decision, renewal, and complaint rules.",
        "Reviewer permissions, assignments, conflicts, scores, comments, decisions, and overrides are auditable.",
        "Notifications are understandable, actionable, role-appropriate, and do not expose sensitive data in subject lines or previews.",
        "The provider has a process for disability accommodations, minors, assisted applications, lost records, and confirmed technical failure.",
        "Benefit delivery, payment failures, in-kind distribution, service delivery, and recipient acknowledgement are traceable.",
        "Retention, deletion, incident handling, and privacy contact responsibilities are documented.",
        "The program is reviewed after each cycle for reach, completion, applicant burden, false mismatches, complaints, and unused criteria.",
    ])

    document.add_heading("42. Common application failures and prevention", level=2)
    add_table(
        document,
        ["Failure point", "Why it happens", "Prevention or fair response"],
        [
            ("Wrong program fit", "A mandatory level, course, location, age, income, or status rule is overlooked.", "Show a plain-language pre-check and criterion-by-criterion explanation before documents."),
            ("Incomplete submission", "A field or required file is missing.", "Use a final checklist and distinguish required from supporting items."),
            ("Unreadable or outdated evidence", "Photo quality, crop, password protection, or validity is poor.", "Preview files, state quality rules, and allow targeted replacement."),
            ("Inconsistent identity or school data", "Names, periods, scales, or institutions differ across records.", "Ask for clarification rather than assuming fraud; preserve the explanation."),
            ("Missed deadline or schedule", "Notice was unclear, contact was inactive, or access was limited.", "Use multiple official notices, visible deadlines, and a documented exception policy."),
            ("False automated mismatch", "Any/custom values, synonyms, missing optional fields, or grading scales are handled badly.", "Use transparent matching, reference normalization, and human review."),
            ("Duplicate or incompatible aid", "The applicant did not understand overlap rules.", "Ask early, name incompatible benefits, and allow informed choice where policy permits."),
            ("No available slot", "More eligible applicants exist than funded places.", "Use ranking or another disclosed allocation method and a controlled waitlist."),
            ("Renewal failure", "Grades, enrollment, units, outputs, or reporting changed.", "Give advance reminders, clear thresholds, correction/appeal routes, and hardship review where allowed."),
        ],
        widths=[Inches(1.45), Inches(2.35), Inches(2.95)],
        font_size=8.8,
    )

    document.add_heading("43. Glossary", level=2)
    glossary = [
        ("Accommodation", "A reasonable adjustment that enables an applicant to access a form, assessment, venue, communication, or program activity."),
        ("Applicant", "The learner seeking support; for a young minor, a guardian may submit on the learner's behalf without becoming the beneficiary."),
        ("Award", "The approved package of support offered under stated conditions."),
        ("Benefit", "Cash, payment, goods, services, mentoring, training, or another item delivered by the program."),
        ("Conditional offer", "An offer that becomes final only when named conditions, such as enrollment, originals, or acceptance, are completed."),
        ("Decision support system", "Rules or analytical tools that organize evidence and guide matching or review while leaving accountable decisions with authorized humans."),
        ("Eligibility", "The mandatory minimum conditions for consideration."),
        ("Formal application", "A provider's official application or due-diligence stage after initial discovery or pre-screening."),
        ("Grant-in-aid", "Educational assistance commonly directed by financial need or a defined expense."),
        ("Guardian-supported account", "An account arrangement in which an authorized adult manages or assists a learner's application while learner and adult data remain distinguishable."),
        ("Match score", "An explanation-based estimate of how well profile information corresponds with program criteria; it is not a final provider score or award."),
        ("Pre-screening", "An initial review of basic eligibility and readiness before more costly, sensitive, or provider-specific steps."),
        ("Provider", "The authorized organization or person that funds, administers, reviews, or delivers the scholarship."),
        ("Required document", "Evidence that must be submitted for the current stage."),
        ("Supporting document", "Additional relevant evidence that may strengthen or clarify an application but is not mandatory unless requested."),
        ("Rubric", "A structured scoring guide with criteria, weights or points, and performance descriptions."),
        ("Scholarship", "Educational support awarded under defined purpose, eligibility, selection, benefit, and continuing conditions."),
        ("Shortlist", "A reduced group still under consideration for a later selection stage."),
        ("Subsidy", "A contribution toward an eligible educational service or cost, often governed separately from scholarships."),
        ("Waitlist", "An ordered or controlled group of eligible applicants who may receive a slot if one becomes available."),
    ]
    for term, definition in glossary:
        add_labeled_paragraph(document, term, definition)

    document.add_heading("44. Official Philippine references and examples", level=2)
    document.add_paragraph(
        "These sources support the distinctions and operational examples in this guide. Program rules, forms, dates, and eligibility may change; always verify the current official issuance or application portal."
    )
    references = [
        ("Republic Act No. 10931, Universal Access to Quality Tertiary Education Act", "https://lawphil.net/statutes/repacts/ra2017/ra_10931_2017.html"),
        ("CHED UniFAST overview", "https://legacy.ched.gov.ph/unifast/"),
        ("CHED Citizen's Charter 2025, including Student Financial Assistance Programs application service", "https://ched.gov.ph/wp-content/uploads/CHED-Updated-CC-2025-1st-edition-033125.pdf"),
        ("CHED Memorandum Order No. 13, series of 2025, revised CHED Merit Scholarship Program guidelines", "https://ched.gov.ph/wp-content/uploads/CMO-No.-13-S.-2025-Revised-CMSP-guidelines.pdf"),
        ("DOST-SEI 2026 Undergraduate Scholarship brochure", "https://science-scholarships.ph/pdf/2026_UG_Scholarship_Brochure.pdf"),
        ("TESDA scholarship programs overview", "https://tesda.gov.ph/About/TESDA/1279"),
        ("TESDA online scholarship application information", "https://www.tesda.gov.ph/Media/NewsDetail/15581"),
        ("DepEd Order No. 020, series of 2023, Guidelines on the Implementation of the Senior High School Voucher Program", "https://deped.gov.ph/wp-content/uploads/DO_s2023_020.pdf"),
        ("DepEd Memorandum No. 030, series of 2025, SHS Voucher Program application timeline", "https://www.deped.gov.ph/wp-content/uploads/DM_s2025_030.pdf"),
        ("PEAC Online Voucher Application Portal eligibility information", "https://ovap.peac.org.ph/registration/eligible"),
        ("National Privacy Commission, Data Privacy Act of 2012 overview", "https://privacy.gov.ph/data-privacy-act/"),
        ("National Privacy Commission Education Sector Advisory No. 2020-1", "https://www.privacy.gov.ph/wp-content/uploads/2020/10/DP-Council-Education-Sector-Advisory-No.-2020-1.pdf"),
        ("National Privacy Commission Advisory Opinion No. 2025-017 concerning minor learner records", "https://privacy.gov.ph/wp-content/uploads/2026/01/NPC-Advisory-Opinion-No.-2025-017_Redactedv2.pdf"),
    ]
    for label, url in references:
        paragraph = document.add_paragraph(style="List Bullet")
        add_hyperlink(paragraph, label, url)

    document.add_heading("45. Final guidance", level=2)
    document.add_paragraph(
        "A trustworthy scholarship process is specific about what support is offered, who it is meant for, why each item of data is needed, how decisions are made, and what happens after selection. The process should be simple where the decision is simple and more detailed only where risk, competition, public accountability, or the learner's pathway requires it."
    )
    document.add_paragraph(
        "For applicants, the safest approach is to verify the source, check mandatory eligibility before uploading documents, submit truthful and readable evidence, monitor official notices, and understand that pre-screening is not the final award. For providers, the strongest approach is to design criteria around a legitimate goal, minimize applicant burden, document decisions, protect data, communicate clear next actions, and evaluate whether the program reaches and supports the learners it intended to serve."
    )
    add_callout(
        document,
        "Remember",
        "The official program announcement and current provider policy determine the actual requirements. When a rule is unclear, applicants should request written clarification through the verified contact before submitting sensitive information or making a financial commitment.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
